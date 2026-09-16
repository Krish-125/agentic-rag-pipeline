from langchain_core.documents import Document
from docx import Document as DocxDocument
import pdfplumber
import pandas as pd
import json
import os

# Load PDF file and extract text and tables into Document objects using PDFPlumber. 
# Each page's text and tables are stored as separate Document objects with metadata 
# including the source filename, page number, and content type (text or table).

def load_pdf(filepath:str) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []

    with pdfplumber.open(filepath) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text()
            if text and text.strip():
                metadata = {
                    "source_filename": filename,
                    "page_number": page_number,
                    "content_type": "text"
                }
                documents.append(Document(page_content=text, metadata=metadata))

            tables = page.extract_tables()
            for t_idx, table in enumerate(tables):
                if not table or len(table) < 2:
                    continue
                headers = table[0]
                table_text = f"Table {t_idx + 1} on page {page_number}:\n"
                for row in table[1:]:
                    row_description = ", ".join(
                        f"{headers[i]} = {row[i]}" for i in range(len(headers)) if i < len(row) and row[i]
                    )
                    table_text += f"- {row_description}\n"
                metadata = {
                    "source_filename": filename,
                    "page_number": page_number,
                    "content_type": "table"
                }
                documents.append(Document(page_content=table_text, metadata=metadata))
    return documents

# Load DOCX file and extract text into Document objects using python-docx.
def load_docx(filepath:str) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []
    doc = DocxDocument(filepath)

    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if full_text.strip():
        metadata = {
            "source_filename": filename,
            "content_type": "text"
        }
        documents.append(Document(page_content=full_text, metadata=metadata))

    for t_idx, table in enumerate(doc.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) < 2:
            continue

        headers = rows[0]
        table_text = f"Table {t_idx + 1}:\n"
        for row in rows[1:]:
            row_description = ", ".join(
                f"{headers[i]} = {row[i]}" for i in range(len(headers)) if i < len(row) and row[i]
            )
            table_text += f"- {row_description}\n"
        metadata = {
            "source_filename": filename,
            "content_type": "table"
        }
        documents.append(Document(page_content=table_text, metadata=metadata))

    return documents

# Load CSV file and extract text into Document objects using the pandas library.
# Chunks the CSV into smaller parts based on the specified number of rows per chunk.
def load_csv(filepath:str, rows_per_chunk:int=20) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []

    df = pd.read_csv(filepath)
    if df.empty:
        return documents

    headers = df.columns.tolist()
    rows = df.values.tolist()

    for chunk_start in range(0, len(rows), rows_per_chunk):
        chunk_rows = rows[chunk_start:chunk_start+rows_per_chunk]
        table_text = f"CSV Table (rows {chunk_start + 1} - {chunk_start + len(chunk_rows)}):\n"  

        for row in chunk_rows:
            row_description = ", ".join(
                f"{headers[i]} = {row[i]}" for i in range(len(headers)) 
                if i < len(row) and pd.notna(row[i]) and row[i] != "" 
            )
            table_text += f"- {row_description}\n"
        metadata = {
            "source_filename": filename,
            "content_type": "table",
            "row_range": f"{chunk_start+1}-{chunk_start+len(chunk_rows)}"
        }
        documents.append(Document(page_content=table_text, metadata=metadata))

    return documents

# Load JSON file and extract text into Document objects using the json library.
# Chunks the JSON into smaller parts based on the specified number of rows per chunk.
def flatten_json(obj, parent_key="") -> list[str]:
    """Recursively turns nested dict/list structures into flat 'key: value' lines."""
    lines = []

    if isinstance(obj, dict):
        for key, value in obj.items():
            full_key = f"{parent_key}.{key}" if parent_key else key
            lines.extend(flatten_json(value, full_key))
    elif isinstance(obj, list):
        for idx, value in enumerate(obj):
            full_key = f"{parent_key}[{idx}]"
            lines.extend(flatten_json(value, full_key))
    else:
        lines.append(f"{parent_key}: {obj}")

    return lines

def load_json(filepath:str, rows_per_chunk:int=25) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []

    with open(filepath, 'r', encoding='utf-8') as f:
        json_data = json.load(f)

    if isinstance(json_data, dict):
        json_data = [json_data]

    for chunk_start in range(0, len(json_data), rows_per_chunk):
        chunk_items = json_data[chunk_start:chunk_start + rows_per_chunk]
        chunk_text = f"JSON Object (items {chunk_start + 1} - {chunk_start + len(chunk_items)}):"

        for i,item in enumerate(chunk_items):
            item_number = chunk_start + i + 1
            chunk_text += f"\nItem {item_number}:"

            if isinstance(item, dict):
                lines = flatten_json(item)
                chunk_text += "\n" + "\n".join(f"- {line}" for line in lines)
            else:
                chunk_text += f"\n- {item}"

        metadata = {
            "source_filename": filename,
            "content_type": "json",
            "item_range": f"{chunk_start+1}-{chunk_start+len(chunk_items)}"
        }
        documents.append(Document(page_content=chunk_text, metadata=metadata))

    return documents

# Load TXT file and extract text into Document objects.
def load_txt(filepath:str) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []

    with open(filepath, 'r', encoding='utf-8') as f:
        text = f.read()

    if text.strip():
        metadata = {
            "source_filename": filename,
            "content_type": "text"
        }
        documents.append(Document(page_content=text, metadata=metadata))

    return documents

# Load Markdown file and extract text into Document objects.
def load_md(filepath:str) -> list[Document]:
    filename = os.path.basename(filepath)
    documents = []

    with open(filepath, 'r', encoding='utf-8') as f:
        text = f.read()

    if text.strip():
        metadata = {
            "source_filename": filename,
            "content_type": "markdown"
        }
        documents.append(Document(page_content=text, metadata=metadata))

    return documents


# Load a document based on its file extension and return a list of Document objects.
def load_document(filepath:str, rows_per_chunk:int=20, json_rows_per_chunk: int = 25) -> list[Document]:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return load_pdf(filepath)
    elif ext == ".docx":
        return load_docx(filepath)
    elif ext == ".csv":
        return load_csv(filepath, rows_per_chunk)
    elif ext == ".json":
        return load_json(filepath, json_rows_per_chunk)
    elif ext == ".txt":
        return load_txt(filepath)
    elif ext == ".md":
        return load_md(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")